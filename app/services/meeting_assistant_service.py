"""Consent-controlled microphone recording and AI meeting-note generation."""

from __future__ import annotations

import os
import re
import io
import wave
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

try:
    from dotenv import load_dotenv

    load_dotenv()
except ImportError:
    pass


@dataclass(frozen=True)
class MeetingReportResult:
    transcript: str
    summary: str
    report_path: Path


class MeetingAssistantService:
    """Transcribe in-memory audio and retain only a written meeting report."""

    SAMPLE_RATE = 16_000

    def __init__(self) -> None:
        self._stream = None
        self._chunks: list[bytes] = []

    @property
    def is_recording(self) -> bool:
        return self._stream is not None

    def start_recording(self) -> None:
        """Start microphone capture after the UI has collected explicit consent."""
        if self.is_recording:
            return
        if not os.getenv("OPENAI_API_KEY", "").strip():
            raise RuntimeError("Add OPENAI_API_KEY to .env before using AI Meeting Notes.")
        try:
            import sounddevice as sound_device
        except ImportError as error:
            raise RuntimeError(
                "Microphone support is missing. Run: pip install -r requirements.txt"
            ) from error

        self._chunks = []

        def receive_audio(input_data, _frames, _time, status) -> None:
            if status:
                print(f"Meeting recorder status: {status}")
            self._chunks.append(bytes(input_data))

        try:
            self._stream = sound_device.RawInputStream(
                samplerate=self.SAMPLE_RATE,
                channels=1,
                dtype="int16",
                callback=receive_audio,
            )
            self._stream.start()
        except Exception as error:
            self._stream = None
            self._chunks = []
            raise RuntimeError(f"The microphone could not be started: {error}") from error

    def stop_and_generate_report(
        self,
        title: str,
        event_date: str,
        start_time: str = "",
        location: str = "",
    ) -> MeetingReportResult:
        """Create a transcript and Word report without writing an audio file."""
        if not self.is_recording:
            raise RuntimeError("No meeting recording is active.")
        self._stop_stream()
        audio = b"".join(self._chunks)
        self._chunks = []
        if not audio:
            raise RuntimeError("No audio was captured. Check the microphone and try again.")

        audio_buffer = io.BytesIO()
        try:
            with wave.open(audio_buffer, "wb") as wav_file:
                wav_file.setnchannels(1)
                wav_file.setsampwidth(2)
                wav_file.setframerate(self.SAMPLE_RATE)
                wav_file.writeframes(audio)
            transcript, summary = self._transcribe_and_summarize(
                audio_buffer.getvalue()
            )
        finally:
            # The microphone data exists only in memory while transcription runs.
            audio_buffer.close()
            del audio
        report_path = self._create_word_report(
            title, event_date, start_time, location, transcript, summary
        )
        return MeetingReportResult(transcript, summary, report_path)

    def cancel_recording(self) -> None:
        """Discard an active recording without uploading or saving it."""
        self._stop_stream()
        self._chunks = []

    def _stop_stream(self) -> None:
        stream = self._stream
        self._stream = None
        if stream is None:
            return
        try:
            stream.stop()
        finally:
            stream.close()

    @staticmethod
    def _transcribe_and_summarize(audio_bytes: bytes) -> tuple[str, str]:
        try:
            from openai import OpenAI
        except ImportError as error:
            raise RuntimeError(
                "OpenAI support is missing. Run: pip install -r requirements.txt"
            ) from error

        client = OpenAI(api_key=os.getenv("OPENAI_API_KEY", "").strip())
        transcription_model = os.getenv(
            "OPENAI_TRANSCRIPTION_MODEL", "gpt-4o-mini-transcribe"
        ).strip()
        notes_model = os.getenv("OPENAI_MEETING_NOTES_MODEL", "gpt-5.4-mini").strip()

        transcription = client.audio.transcriptions.create(
            model=transcription_model,
            file=("meeting.wav", audio_bytes, "audio/wav"),
        )
        transcript = str(getattr(transcription, "text", "") or "").strip()
        if not transcript:
            raise RuntimeError("The recording could not be transcribed.")

        response = client.responses.create(
            model=notes_model,
            store=False,
            instructions=(
                "Create concise, factual internal meeting notes. Use short sections for "
                "Summary, Decisions, and Action items with owners and deadlines when stated. "
                "Do not invent missing facts. Clearly label anything uncertain."
            ),
            input=transcript,
        )
        notes = str(getattr(response, "output_text", "") or "").strip()
        if not notes:
            raise RuntimeError("AI notes were not returned.")
        return transcript, notes

    @staticmethod
    def _create_word_report(
        title: str,
        event_date: str,
        start_time: str,
        location: str,
        transcript: str,
        summary: str,
        output_directory: Path | None = None,
    ) -> Path:
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt, RGBColor
        except ImportError as error:
            raise RuntimeError(
                "Word report support is missing. Run: pip install -r requirements.txt"
            ) from error

        report_directory = output_directory or (
            Path(__file__).resolve().parents[2] / "data" / "meeting_reports"
        )
        report_directory.mkdir(parents=True, exist_ok=True)
        safe_title = re.sub(r"[^A-Za-z0-9_-]+", "_", title.strip()).strip("_") or "Meeting"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_path = report_directory / f"{safe_title}_{timestamp}.docx"

        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(10.5)
        for style_name in ("Title", "Heading 1", "Heading 2"):
            styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

        heading = document.add_paragraph(style="Title")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run("Meeting Report")
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.add_run(title.strip() or "Untitled meeting").bold = True

        metadata = document.add_table(rows=0, cols=2)
        metadata.style = "Table Grid"
        for label, value in (
            ("Date", event_date or "Not specified"),
            ("Time", start_time or "Not specified"),
            ("Location", location or "Not specified"),
            ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ):
            cells = metadata.add_row().cells
            cells[0].text = label
            cells[1].text = value
            cells[0].paragraphs[0].runs[0].bold = True

        document.add_heading("Meeting Summary", level=1)
        document.add_paragraph(summary)
        document.add_heading("Full Transcript", level=1)
        document.add_paragraph(transcript)
        document.add_paragraph()
        notice = document.add_paragraph(
            "Generated by Untangled Nexus AI Meeting Assistant. Review for accuracy before distribution."
        )
        notice.runs[0].italic = True
        notice.runs[0].font.size = Pt(9)
        notice.runs[0].font.color.rgb = RGBColor(90, 90, 90)
        document.core_properties.title = f"Meeting Report {title.strip()}"
        document.core_properties.author = "Untangled Nexus"
        document.save(report_path)
        return report_path
